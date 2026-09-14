import io
from datetime import date, datetime
from decimal import Decimal

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from ..database import get_db
from ..dependencies import require_roles
from ..models import Usuario, Venta

router = APIRouter(prefix="/api/reportes", tags=["Reportes"])


def _obtener_ventas_del_dia(db: Session, fecha: date) -> list[Venta]:
    return (
        db.query(Venta)
        .options(joinedload(Venta.detalles), joinedload(Venta.cliente))
        .filter(func.date(Venta.creado_en) == fecha)
        .order_by(Venta.creado_en)
        .all()
    )


def _nombre_cliente(venta: Venta) -> str:
    if venta.cliente:
        return f"{venta.cliente.nombre} {venta.cliente.apellido}"
    return "N/A"


def _resumen_items(venta: Venta) -> str:
    return ", ".join(f"{d.nombre_item} x{d.cantidad}" for d in venta.detalles)


@router.get("/ventas/diario")
def reporte_diario_json(
    fecha: date = Query(..., description="Fecha del reporte, formato YYYY-MM-DD"),
    db: Session = Depends(get_db),
    usuario_actual: Usuario = Depends(require_roles("administrador", "empleado")),
):
    """Vista previa en JSON del reporte diario, útil para mostrarlo en el frontend
    antes de descargar el PDF/Excel."""
    ventas = _obtener_ventas_del_dia(db, fecha)
    total_dia = sum((v.total for v in ventas), Decimal("0"))

    filas = [
        {
            "numero_venta": v.id,
            "cliente": _nombre_cliente(v),
            "items": _resumen_items(v),
            "total": float(v.total),
            "estado": v.estado,
        }
        for v in ventas
    ]

    return {
        "fecha": fecha.isoformat(),
        "cantidad_ventas": len(ventas),
        "total_dia": float(total_dia),
        "ventas": filas,
    }


@router.get("/ventas/diario/pdf")
def reporte_diario_pdf(
    fecha: date = Query(...),
    db: Session = Depends(get_db),
    usuario_actual: Usuario = Depends(require_roles("administrador", "empleado")),
):
    ventas = _obtener_ventas_del_dia(db, fecha)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, title=f"Reporte de ventas {fecha}")
    estilos = getSampleStyleSheet()
    elementos = []

    elementos.append(Paragraph("Sweet Ice - Reporte Diario de Ventas", estilos["Title"]))
    elementos.append(Paragraph(f"Fecha: {fecha.strftime('%d/%m/%Y')}", estilos["Normal"]))
    elementos.append(Spacer(1, 14))

    datos_tabla = [["N° Venta", "Cliente", "Productos/Servicios", "Total", "Estado"]]
    total_dia = Decimal("0")
    for v in ventas:
        datos_tabla.append(
            [str(v.id), _nombre_cliente(v), _resumen_items(v), f"${v.total:,.0f}", v.estado]
        )
        total_dia += v.total

    tabla = Table(datos_tabla, colWidths=[55, 100, 210, 70, 70], repeatRows=1)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#db2777")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fff1f2")]),
            ]
        )
    )
    elementos.append(tabla)
    elementos.append(Spacer(1, 18))
    elementos.append(Paragraph(f"<b>Cantidad de ventas:</b> {len(ventas)}", estilos["Normal"]))
    elementos.append(Paragraph(f"<b>Total del día:</b> ${total_dia:,.0f}", estilos["Normal"]))
    elementos.append(Spacer(1, 24))
    elementos.append(
        Paragraph(
            f"Reporte generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            estilos["Italic"],
        )
    )

    doc.build(elementos)
    buffer.seek(0)

    nombre_archivo = f"reporte_ventas_{fecha.isoformat()}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )


@router.get("/ventas/diario/excel")
def reporte_diario_excel(
    fecha: date = Query(...),
    db: Session = Depends(get_db),
    usuario_actual: Usuario = Depends(require_roles("administrador", "empleado")),
):
    ventas = _obtener_ventas_del_dia(db, fecha)

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte Ventas"

    ws.merge_cells("A1:F1")
    ws["A1"] = f"Sweet Ice - Reporte Diario de Ventas - {fecha.strftime('%d/%m/%Y')}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.append([])
    ws.append(["N° Venta", "Cliente", "Productos/Servicios", "Cant. items", "Total", "Estado"])
    fila_encabezado = ws.max_row
    for celda in ws[fila_encabezado]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(start_color="DB2777", end_color="DB2777", fill_type="solid")
        celda.alignment = Alignment(horizontal="center")

    total_dia = Decimal("0")
    for v in ventas:
        cantidad_items = sum(d.cantidad for d in v.detalles)
        ws.append(
            [v.id, _nombre_cliente(v), _resumen_items(v), cantidad_items, float(v.total), v.estado]
        )
        total_dia += v.total

    ws.append([])
    fila_total = ws.max_row + 1
    ws[f"A{fila_total}"] = "Total del día:"
    ws[f"A{fila_total}"].font = Font(bold=True)
    ws[f"E{fila_total}"] = float(total_dia)
    ws[f"E{fila_total}"].font = Font(bold=True)

    fila_meta = fila_total + 2
    ws[f"A{fila_meta}"] = f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}"

    anchos = {"A": 10, "B": 25, "C": 45, "D": 13, "E": 15, "F": 12}
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"reporte_ventas_{fecha.isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )